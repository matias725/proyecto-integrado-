"""Exportación a Excel de los resultados filtrados (todas las páginas)."""

from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROJO_INSTITUCIONAL = '920000'


def respuesta_excel(nombre_archivo, titulo, encabezados, filas, filtros, usuario):
    libro = Workbook()
    hoja = libro.active
    hoja.title = titulo[:31]

    hoja['A1'] = titulo
    hoja['A1'].font = Font(bold=True, size=14, color=ROJO_INSTITUCIONAL)
    ahora = timezone.localtime()
    hoja['A2'] = f'Generado el {ahora:%d-%m-%Y %H:%M} por {usuario.nombre}'
    hoja['A3'] = 'Filtros: ' + ('; '.join(f'{k}: {v}' for k, v in filtros) or 'sin filtros')
    hoja['A4'] = f'Total de registros: {len(filas)}'
    for celda in ('A2', 'A3', 'A4'):
        hoja[celda].font = Font(italic=True, color='5C554E')

    inicio = 6
    for col, encabezado in enumerate(encabezados, start=1):
        celda = hoja.cell(row=inicio, column=col, value=encabezado)
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor=ROJO_INSTITUCIONAL)
        celda.alignment = Alignment(vertical='center')

    anchos = [len(str(e)) for e in encabezados]
    for fila_n, fila in enumerate(filas, start=inicio + 1):
        for col, valor in enumerate(fila, start=1):
            if hasattr(valor, 'tzinfo') and valor.tzinfo is not None:
                valor = timezone.localtime(valor).replace(tzinfo=None)
            celda = hoja.cell(row=fila_n, column=col, value=valor)
            if isinstance(valor, str) and valor.startswith('='):
                celda.data_type = 's'  # texto literal: evita inyección de fórmulas
            anchos[col - 1] = max(anchos[col - 1], len(str(valor or '')))

    for col, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(col)].width = min(max(ancho + 2, 10), 60)
    hoja.freeze_panes = hoja.cell(row=inicio + 1, column=1)
    if filas:
        hoja.auto_filter.ref = f'A{inicio}:{get_column_letter(len(encabezados))}{inicio + len(filas)}'

    salida = BytesIO()
    libro.save(salida)
    respuesta = HttpResponse(
        salida.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    respuesta['Content-Disposition'] = f'attachment; filename="{nombre_archivo}-{ahora:%Y%m%d-%H%M}.xlsx"'
    return respuesta
