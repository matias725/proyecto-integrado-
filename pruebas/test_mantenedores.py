"""Pruebas de los seis mantenedores (CRUD), listados y exportación a Excel."""

from io import BytesIO

from django.test import override_settings
from django.urls import reverse
from openpyxl import load_workbook

from atenciones.models import Atencion
from catalogos.models import Delegacion, ItemMedicion, SubAtencion, TipoAtencion
from core.models import Auditoria
from cuentas.models import Permiso, Rol, Usuario

from .base import PruebaSGR


def leer_excel(respuesta):
    hoja = load_workbook(BytesIO(respuesta.content)).active
    filas = list(hoja.iter_rows(values_only=True))
    encabezados = filas[5]
    datos = [f for f in filas[6:] if any(c is not None for c in f)]
    return hoja, encabezados, datos


class CrudGenericoTests(PruebaSGR):
    """Crear, consultar, listar, editar, desactivar y eliminar en cada mantenedor."""

    CASOS = [
        # modulo, modelo, datos de alta, campo a editar
        ('delegaciones', Delegacion,
         {'codigo': 'DEL07', 'nombre': 'Serena Oriente', 'direccion': 'Calle Prueba 123', 'comuna': 'La Serena',
          'activo': 'on'}, ('direccion', 'Calle Modificada 456')),
        ('items', ItemMedicion,
         {'codigo': 'IT099', 'nombre': 'Ítem de prueba', 'descripcion': 'Descripción del ítem', 'activo': 'on'},
         ('descripcion', 'Descripción editada')),
        ('tipos_atencion', TipoAtencion,
         {'codigo': 'TA099', 'nombre': 'Tipo de prueba', 'descripcion': 'Descripción del tipo', 'activo': 'on'},
         ('nombre', 'Tipo de prueba editado')),
        ('sub_atenciones', SubAtencion,
         {'codigo': 'SA099', 'nombre': 'Sub de prueba', 'tipo': None, 'activo': 'on'},
         ('nombre', 'Sub de prueba editada')),
        ('roles', Rol,
         {'codigo': 'ROL09', 'nombre': 'Rol de prueba', 'descripcion': 'Rol creado en pruebas',
          'alcance': Rol.DELEGACION, 'permisos': None, 'activo': 'on'},
         ('descripcion', 'Descripción editada del rol')),
        ('usuarios', Usuario,
         {'nombre': 'Persona de Prueba', 'nombre_usuario': 'pprueba', 'correo': 'pprueba@example.test',
          'rol': None, 'delegaciones': None, 'is_active': 'on', 'clave': 'Elqui#Valle77', 'confirmacion': 'Elqui#Valle77'},
         ('nombre', 'Persona Editada')),
    ]

    def preparar(self, modulo, datos):
        datos = dict(datos)
        if modulo == 'sub_atenciones':
            datos['tipo'] = TipoAtencion.objects.get(codigo='TA001').pk
        if modulo == 'roles':
            datos['permisos'] = list(Permiso.objects.filter(modulo='atenciones').values_list('pk', flat=True))
        if modulo == 'usuarios':
            datos['rol'] = Rol.objects.get(codigo='ROL04').pk
            datos['delegaciones'] = [Delegacion.objects.get(codigo='DEL02').pk]
        return datos

    def test_ciclo_completo_en_los_seis_mantenedores(self):
        """PC-01 · Crear, consultar, listar, editar, desactivar y eliminar en los 6 mantenedores"""
        self.entrar('admin')
        for modulo, modelo, datos, (campo, nuevo_valor) in self.CASOS:
            with self.subTest(modulo=modulo):
                datos = self.preparar(modulo, datos)
                # Crear
                r = self.client.post(reverse(f'{modulo}_crear'), datos)
                self.assertEqual(r.status_code, 302, getattr(r, 'context', None) and r.context['form'].errors)
                clave = 'nombre_usuario' if modulo == 'usuarios' else 'codigo'
                obj = modelo.objects.get(**{clave: datos[clave]})
                self.assertRedirects(r, reverse(f'{modulo}_ver', args=[obj.pk]))
                self.assertTrue(Auditoria.objects.filter(evento='ALTA', entidad=modulo, registro=str(obj.pk)).exists())
                # Consultar
                r = self.client.get(reverse(f'{modulo}_ver', args=[obj.pk]))
                self.assertEqual(r.status_code, 200)
                # Listar y buscar
                r = self.client.get(reverse(f'{modulo}_lista'), {'q': datos.get('nombre', '')[:8]})
                self.assertContains(r, datos['nombre'][:8])
                # Editar
                datos_edit = dict(datos, **{campo: nuevo_valor})
                if modulo == 'usuarios':
                    datos_edit.update(clave='', confirmacion='')
                r = self.client.post(reverse(f'{modulo}_editar', args=[obj.pk]), datos_edit)
                self.assertEqual(r.status_code, 302)
                obj.refresh_from_db()
                self.assertEqual(getattr(obj, campo), nuevo_valor)
                # Desactivar y reactivar
                activo = 'is_active' if modulo == 'usuarios' else 'activo'
                self.client.post(reverse(f'{modulo}_estado', args=[obj.pk]))
                obj.refresh_from_db()
                self.assertFalse(getattr(obj, activo))
                self.client.post(reverse(f'{modulo}_estado', args=[obj.pk]))
                obj.refresh_from_db()
                self.assertTrue(getattr(obj, activo))
                # Eliminar (sin relaciones)
                r = self.client.post(reverse(f'{modulo}_eliminar', args=[obj.pk]))
                self.assertRedirects(r, reverse(f'{modulo}_lista'))
                self.assertFalse(modelo.objects.filter(pk=obj.pk).exists())
                self.assertTrue(Auditoria.objects.filter(evento='ELIMINACION', entidad=modulo).exists())

    def test_eliminar_exige_post(self):
        """PC-02 · Eliminar y cambiar estado no se ejecutan por GET"""
        self.entrar('admin')
        item = ItemMedicion.objects.first()
        self.assertEqual(self.client.get(reverse('items_eliminar', args=[item.pk])).status_code, 405)
        self.assertEqual(self.client.get(reverse('items_estado', args=[item.pk])).status_code, 405)
        self.assertTrue(ItemMedicion.objects.filter(pk=item.pk).exists())

    def test_campos_obligatorios_en_servidor(self):
        """PC-03 · Campos obligatorios validados en el servidor"""
        self.entrar('admin')
        for modulo in ['delegaciones', 'items', 'tipos_atencion', 'sub_atenciones', 'roles', 'usuarios']:
            with self.subTest(modulo=modulo):
                r = self.client.post(reverse(f'{modulo}_crear'), {})
                self.assertEqual(r.status_code, 200)
                self.assertTrue(r.context['form'].errors)
                self.assertMensaje(r, 'No se guardó')

    def test_codigo_y_nombre_duplicados(self):
        """PC-04 · Código o nombre duplicado (sin distinguir mayúsculas) es rechazado"""
        self.entrar('admin')
        r = self.client.post(reverse('delegaciones_crear'), {
            'codigo': 'del01', 'nombre': 'centro', 'direccion': 'X 1', 'comuna': 'La Serena', 'activo': 'on'})
        self.assertEqual(r.status_code, 200)
        self.assertIn('codigo', r.context['form'].errors)
        self.assertIn('nombre', r.context['form'].errors)

    def test_formato_de_codigo(self):
        """PC-05 · El código debe respetar el formato del catálogo (p. ej. TA001)"""
        self.entrar('admin')
        r = self.client.post(reverse('tipos_atencion_crear'), {'codigo': 'X1', 'nombre': 'Nuevo', 'descripcion': 'D'})
        self.assertIn('codigo', r.context['form'].errors)

    def test_codigo_propuesto(self):
        """PC-06 · El formulario de alta propone el siguiente código disponible"""
        self.entrar('admin')
        r = self.client.get(reverse('delegaciones_crear'))
        self.assertEqual(r.context['form']['codigo'].value(), 'DEL07')


class IntegridadTests(PruebaSGR):
    """Reglas que protegen las relaciones entre registros."""

    def setUp(self):
        self.entrar('admin')

    def test_no_eliminar_delegacion_con_usuarios(self):
        """PI-01 · No se elimina una delegación con usuarios o atenciones"""
        d = Delegacion.objects.get(codigo='DEL06')
        r = self.client.post(reverse('delegaciones_eliminar', args=[d.pk]), follow=True)
        self.assertTrue(Delegacion.objects.filter(pk=d.pk).exists())
        self.assertMensaje(r, 'No se pudo eliminar')

    def test_no_desactivar_delegacion_con_usuarios_activos(self):
        """PI-02 · No se desactiva una delegación con usuarios activos"""
        d = Delegacion.objects.get(codigo='DEL06')
        r = self.client.post(reverse('delegaciones_estado', args=[d.pk]), follow=True)
        d.refresh_from_db()
        self.assertTrue(d.activo)
        self.assertMensaje(r, 'No se pudo desactivar')

    def test_no_eliminar_tipo_con_subatenciones(self):
        """PI-03 · No se elimina un tipo de atención con sub atenciones"""
        t = TipoAtencion.objects.get(codigo='TA001')
        self.client.post(reverse('tipos_atencion_eliminar', args=[t.pk]))
        self.assertTrue(TipoAtencion.objects.filter(pk=t.pk).exists())

    def test_no_desactivar_tipo_con_subatenciones_activas(self):
        """PI-04 · No se desactiva un tipo con sub atenciones activas"""
        t = TipoAtencion.objects.get(codigo='TA001')
        r = self.client.post(reverse('tipos_atencion_estado', args=[t.pk]), follow=True)
        t.refresh_from_db()
        self.assertTrue(t.activo)
        self.assertMensaje(r, 'sub atenciones activas')

    def test_no_activar_sub_con_tipo_inactivo(self):
        """PI-05 · Una sub atención no puede activarse si su tipo está inactivo"""
        tipo = TipoAtencion.objects.get(codigo='TA009')
        sub = SubAtencion.objects.create(codigo='SA098', nombre='Sub huérfana', tipo=tipo, activo=False)
        r = self.client.post(reverse('sub_atenciones_estado', args=[sub.pk]), follow=True)
        sub.refresh_from_db()
        self.assertFalse(sub.activo)
        self.assertMensaje(r, 'No se pudo activar')

    def test_sub_solo_con_tipos_activos(self):
        """PI-06 · Al crear una sub atención solo se ofrecen tipos activos"""
        r = self.client.get(reverse('sub_atenciones_crear'))
        tipos = r.context['form'].fields['tipo'].queryset
        self.assertFalse(tipos.filter(activo=False).exists())
        inactivo = TipoAtencion.objects.get(codigo='TA009')
        r = self.client.post(reverse('sub_atenciones_crear'),
                             {'codigo': 'SA097', 'nombre': 'Otra', 'tipo': inactivo.pk, 'activo': 'on'})
        self.assertIn('tipo', r.context['form'].errors)

    def test_nombre_sub_unico_dentro_del_tipo(self):
        """PI-07 · El nombre de sub atención es único dentro de su tipo"""
        existente = SubAtencion.objects.filter(activo=True).first()
        r = self.client.post(reverse('sub_atenciones_crear'), {
            'codigo': 'SA096', 'nombre': existente.nombre.upper(), 'tipo': existente.tipo_id, 'activo': 'on'})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.context['form'].errors)

    def test_no_eliminar_sub_con_atenciones(self):
        """PI-08 · No se elimina una sub atención usada en atenciones"""
        sub = Atencion.objects.first().subatencion
        self.client.post(reverse('sub_atenciones_eliminar', args=[sub.pk]))
        self.assertTrue(SubAtencion.objects.filter(pk=sub.pk).exists())

    def test_rol_protegido(self):
        """PI-09 · El rol Superadministrador no se elimina ni se desactiva"""
        rol = Rol.objects.get(codigo='ROL01')
        self.client.post(reverse('roles_eliminar', args=[rol.pk]))
        self.client.post(reverse('roles_estado', args=[rol.pk]))
        rol.refresh_from_db()
        self.assertTrue(rol.activo)

    def test_no_eliminar_rol_con_usuarios(self):
        """PI-10 · No se elimina un rol asignado a usuarios"""
        rol = Rol.objects.get(codigo='ROL04')
        self.client.post(reverse('roles_eliminar', args=[rol.pk]))
        self.assertTrue(Rol.objects.filter(pk=rol.pk).exists())

    def test_no_eliminarse_ni_desactivarse_a_si_mismo(self):
        """PI-11 · Un usuario no puede eliminarse ni desactivarse a sí mismo"""
        yo = self.usuario('admin')
        self.client.post(reverse('usuarios_eliminar', args=[yo.pk]))
        self.client.post(reverse('usuarios_estado', args=[yo.pk]))
        yo.refresh_from_db()
        self.assertTrue(yo.is_active)

    def test_no_eliminar_usuario_con_atenciones(self):
        """PI-12 · No se elimina un usuario con atenciones registradas (se desactiva)"""
        u = self.usuario('bcastillo')
        self.client.post(reverse('usuarios_eliminar', args=[u.pk]))
        self.assertTrue(Usuario.objects.filter(pk=u.pk).exists())

    def test_usuario_correo_y_usuario_unicos(self):
        """PI-13 · Correo y nombre de usuario únicos"""
        r = self.client.post(reverse('usuarios_crear'), {
            'nombre': 'Copia', 'nombre_usuario': 'BCASTILLO', 'correo': 'BCastillo@example.test',
            'rol': Rol.objects.get(codigo='ROL04').pk, 'delegaciones': [Delegacion.objects.first().pk],
            'is_active': 'on', 'clave': 'Elqui#Valle77', 'confirmacion': 'Elqui#Valle77'})
        self.assertIn('nombre_usuario', r.context['form'].errors)
        self.assertIn('correo', r.context['form'].errors)

    def test_usuario_clave_debil(self):
        """PI-14 · Al crear un usuario la contraseña debe cumplir los requisitos"""
        r = self.client.post(reverse('usuarios_crear'), {
            'nombre': 'Nuevo', 'nombre_usuario': 'nuevo1', 'correo': 'nuevo1@example.test',
            'rol': Rol.objects.get(codigo='ROL04').pk, 'delegaciones': [Delegacion.objects.first().pk],
            'is_active': 'on', 'clave': 'abc', 'confirmacion': 'abc'})
        self.assertIn('clave', r.context['form'].errors)

    def test_rol_desactivado_no_puede_asignarse(self):
        """PI-15 · No se ofrece un rol inactivo al crear usuarios"""
        rol = Rol.objects.create(codigo='ROL08', nombre='Rol inactivo', descripcion='x', activo=False)
        r = self.client.get(reverse('usuarios_crear'))
        self.assertNotIn(rol, r.context['form'].fields['rol'].queryset)


@override_settings(SGR_POR_PAGINA=5)
class ListadosTests(PruebaSGR):
    """Paginador, total, búsqueda, filtros y exportación coherente."""

    def setUp(self):
        self.entrar('admin')

    def test_paginador_y_total(self):
        """PL-01 · El listado pagina y muestra el total de registros"""
        r = self.client.get(reverse('items_lista'))
        self.assertEqual(r.context['pagina'].paginator.count, 14)
        self.assertEqual(len(r.context['filas']), 5)
        self.assertContains(r, 'de 14 registros')
        r = self.client.get(reverse('items_lista'), {'pagina': 3})
        self.assertEqual(len(r.context['filas']), 4)

    def test_pagina_fuera_de_rango(self):
        """PL-02 · Una página inexistente muestra la última válida (sin error)"""
        r = self.client.get(reverse('items_lista'), {'pagina': 999})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.context['pagina'].number, 3)
        r = self.client.get(reverse('items_lista'), {'pagina': 'abc'})
        self.assertEqual(r.context['pagina'].number, 1)

    def test_busqueda_parcial_multicampo_sin_mayusculas(self):
        """PL-03 · Búsqueda parcial, en varios campos y sin distinguir mayúsculas"""
        r = self.client.get(reverse('delegaciones_lista'), {'q': 'aNTEN'})
        codigos = [f['obj'].codigo for f in r.context['filas']]
        self.assertEqual(codigos, ['DEL03'])
        r = self.client.get(reverse('delegaciones_lista'), {'q': 'esmeralda'})  # busca en dirección
        self.assertEqual([f['obj'].codigo for f in r.context['filas']], ['DEL04'])
        r = self.client.get(reverse('sub_atenciones_lista'), {'q': 'salud'})  # busca en tipo asociado
        self.assertTrue(r.context['filas'])

    def test_busqueda_varias_palabras(self):
        """PL-04 · Varias palabras se combinan (todas deben coincidir)"""
        r = self.client.get(reverse('usuarios_lista'), {'q': 'rojas funcionario'})
        nombres = {f['obj'].nombre_usuario for f in r.context['filas']}
        self.assertIn('fortiz', nombres)
        self.assertNotIn('arojas', nombres)

    def test_sin_coincidencias(self):
        """PL-05 · Estado «sin coincidencias» distinto del listado vacío"""
        r = self.client.get(reverse('items_lista'), {'q': 'zzzz-no-existe'})
        self.assertContains(r, 'No hay coincidencias')

    def test_filtros_se_conservan_al_paginar(self):
        """PL-06 · Los filtros se conservan en los enlaces del paginador"""
        r = self.client.get(reverse('usuarios_lista'), {'q': 'a', 'estado': 'activos'})
        self.assertContains(r, 'q=a&amp;estado=activos&amp;pagina=2')

    def test_filtro_por_estado(self):
        """PL-07 · Filtro por estado activo/inactivo"""
        r = self.client.get(reverse('tipos_atencion_lista'), {'estado': 'inactivos'})
        self.assertEqual([f['obj'].codigo for f in r.context['filas']], ['TA009'])

    def test_estado_textual_en_tabla(self):
        """PL-08 · La tabla muestra el estado en texto (no solo color)"""
        r = self.client.get(reverse('tipos_atencion_lista'))
        self.assertContains(r, 'Activo')
        self.assertContains(r, 'Inactivo')

    def test_exportar_todos_los_filtrados(self):
        """PL-09 · Excel incluye todos los resultados filtrados, no solo la página visible"""
        r = self.client.get(reverse('items_exportar'))
        self.assertEqual(r['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        _, encabezados, datos = leer_excel(r)
        self.assertEqual(len(datos), 14)  # la página muestra 5
        self.assertIn('Estado', encabezados)

    def test_exportar_coincide_con_listado_filtrado(self):
        """PL-10 · Excel coincide con el listado filtrado (búsqueda + estado)"""
        filtros = {'q': 'a', 'estado': 'activos'}
        lista = self.client.get(reverse('sub_atenciones_lista'), filtros)
        r = self.client.get(reverse('sub_atenciones_exportar'), filtros)
        hoja, _, datos = leer_excel(r)
        self.assertEqual(len(datos), lista.context['pagina'].paginator.count)
        self.assertIn('Búsqueda: a', hoja['A3'].value)
        self.assertIn(f'Total de registros: {len(datos)}', hoja['A4'].value)

    def test_exportar_usuarios_sin_contrasenas(self):
        """PL-11 · La exportación de usuarios no incluye contraseñas ni datos sensibles"""
        r = self.client.get(reverse('usuarios_exportar'))
        _, encabezados, datos = leer_excel(r)
        texto = ' '.join(str(c) for fila in datos for c in fila) + ' '.join(map(str, encabezados))
        self.assertNotIn('md5$', texto)
        self.assertNotIn('pbkdf2', texto)
        self.assertFalse(any('contrase' in str(e).lower() or 'clave' in str(e).lower() for e in encabezados))

    def test_exportacion_queda_auditada(self):
        """PL-12 · Cada exportación queda registrada en la auditoría"""
        self.client.get(reverse('roles_exportar'))
        self.assertTrue(Auditoria.objects.filter(evento='EXPORTACION', entidad='roles').exists())

    def test_excel_sin_formulas_inyectadas(self):
        """PL-13 · Un texto que empieza con «=» se exporta como texto, no como fórmula"""
        ItemMedicion.objects.create(codigo='IT098', nombre='=1+1', descripcion='=HYPERLINK("x")')
        r = self.client.get(reverse('items_exportar'), {'q': 'IT098'})
        hoja, _, _ = leer_excel(r)
        self.assertEqual(hoja['B7'].data_type, 's')
